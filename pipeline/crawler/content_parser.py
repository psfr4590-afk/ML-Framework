"""Format-aware parsing for web-crawled documents.

The crawler must never turn a binary/data document into an empty or metadata-only
Document. This module converts common document/data formats into deterministic
text while preserving record, page, sheet, and slide boundaries where possible.
Unsupported or image-only content fails closed instead of silently producing bad
training data.
"""
from __future__ import annotations

import csv
import io
import json
import mimetypes
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path


class ContentParseError(ValueError):
    """Raised when content cannot be parsed safely and completely enough to ingest."""


@dataclass(frozen=True)
class ParsedContent:
    text: str
    title: str = ""
    content_type: str = "unknown"
    links: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


_EXTENSIONS = {
    ".pdf": "pdf",
    ".csv": "csv",
    ".tsv": "tsv",
    ".json": "json",
    ".jsonl": "jsonl",
    ".ndjson": "jsonl",
    ".xml": "xml",
    ".html": "html",
    ".htm": "html",
    ".md": "text",
    ".markdown": "text",
    ".txt": "text",
    ".rst": "text",
    ".log": "text",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pptx": "pptx",
}

_MIME_TYPES = {
    "application/pdf": "pdf",
    "text/csv": "csv",
    "application/csv": "csv",
    "text/tab-separated-values": "tsv",
    "application/json": "json",
    "application/ld+json": "json",
    "application/x-ndjson": "jsonl",
    "application/xml": "xml",
    "text/xml": "xml",
    "text/html": "html",
    "application/xhtml+xml": "html",
    "text/plain": "text",
    "text/markdown": "text",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
}


def _kind(url: str, content_type: str) -> str:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix in _EXTENSIONS:
        return _EXTENSIONS[suffix]
    mime = content_type.split(";", 1)[0].strip().lower()
    if mime in _MIME_TYPES:
        return _MIME_TYPES[mime]
    guessed = mimetypes.guess_type(url)[0]
    return _MIME_TYPES.get(guessed or "", "")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "utf-32", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ContentParseError("text content could not be decoded")


def _parse_csv(data: bytes, delimiter: str | None = None) -> ParsedContent:
    text = _decode_text(data).replace("\x00", "")
    sample = text[:65536]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        raise ContentParseError("CSV/TSV contains no rows")
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    headers = rows[0]
    if not any(cell.strip() for cell in headers):
        headers = [f"column_{i + 1}" for i in range(width)]
        start = 0
    else:
        headers = [cell.strip() or f"column_{i + 1}" for i, cell in enumerate(headers)]
        start = 1
    records = []
    for row_number, row in enumerate(rows[start:], start=start + 1):
        if not any(cell.strip() for cell in row):
            continue
        fields = [f"{headers[i]}: {row[i]}" for i in range(width)]
        records.append(f"record {row_number}\n" + "\n".join(fields))
    if not records:
        raise ContentParseError("CSV/TSV contains no non-empty data records")
    return ParsedContent(
        text="\n\n".join(records),
        content_type="tabular_data",
        metadata={"format": "tsv" if delimiter == "\t" else "csv", "rows": len(records), "columns": width, "delimiter": delimiter},
    )


def _json_records(data: object) -> list[object]:
    if isinstance(data, list):
        return data
    return [data]


def _parse_json(data: bytes, jsonl: bool = False) -> ParsedContent:
    text = _decode_text(data)
    records = []
    if jsonl:
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise ContentParseError(f"invalid JSONL at line {line_number}: {exc}") from exc
    else:
        try:
            records = _json_records(json.loads(text))
        except json.JSONDecodeError as exc:
            raise ContentParseError(f"invalid JSON: {exc}") from exc
    if not records:
        raise ContentParseError("JSON contains no records")
    rendered = [json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":")) for record in records]
    return ParsedContent(text="\n".join(rendered), content_type="structured_data", metadata={"format": "jsonl" if jsonl else "json", "records": len(records)})


def _parse_pdf(data: bytes) -> ParsedContent:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ContentParseError("PDF parsing requires pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        nonempty = 0
        for number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                nonempty += 1
                pages.append(f"[PAGE {number}]\n{text}")
        if not pages:
            raise ContentParseError("PDF contains no extractable text; scanned/image-only PDFs require an OCR parser")
        return ParsedContent(
            text="\n\n".join(pages),
            content_type="document_pdf",
            metadata={"format": "pdf", "pages": len(reader.pages), "text_pages": nonempty},
        )
    except ContentParseError:
        raise
    except Exception as exc:
        raise ContentParseError(f"PDF parsing failed: {exc}") from exc


def _parse_xml(data: bytes) -> ParsedContent:
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise ContentParseError(f"invalid XML: {exc}") from exc
    text = " ".join(part.strip() for part in root.itertext() if part and part.strip())
    if not text:
        raise ContentParseError("XML contains no text")
    return ParsedContent(text=text, content_type="structured_data", metadata={"format": "xml", "root": root.tag})


def _parse_docx(data: bytes) -> ParsedContent:
    try:
        from docx import Document as WordDocument
    except ImportError as exc:
        raise ContentParseError("DOCX parsing requires python-docx") from exc
    try:
        doc = WordDocument(io.BytesIO(data))
        blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table_number, table in enumerate(doc.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(f"[TABLE {table_number}]\n" + "\n".join(" | ".join(row) for row in rows))
        if not blocks:
            raise ContentParseError("DOCX contains no extractable text")
        return ParsedContent(text="\n\n".join(blocks), content_type="document", metadata={"format": "docx", "tables": len(doc.tables)})
    except ContentParseError:
        raise
    except Exception as exc:
        raise ContentParseError(f"DOCX parsing failed: {exc}") from exc


def _parse_xlsx(data: bytes) -> ParsedContent:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ContentParseError("XLSX parsing requires openpyxl") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks = []
        total_rows = 0
        sheet_count = len(workbook.sheetnames)
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            total_rows += len(rows)
            headers = [str(v).strip() if v is not None else f"column_{i + 1}" for i, v in enumerate(rows[0])]
            records = [
                "\n".join(f"{headers[i]}: {'' if value is None else value}" for i, value in enumerate(row))
                for row in rows[1:]
                if any(value is not None and str(value).strip() for value in row)
            ]
            if records:
                blocks.append(f"[SHEET {sheet.title}]\n" + "\n\n".join(records))
        workbook.close()
        if not blocks:
            raise ContentParseError("XLSX contains no non-empty data rows")
        return ParsedContent(text="\n\n".join(blocks), content_type="tabular_data", metadata={"format": "xlsx", "sheets": sheet_count, "rows": total_rows})
    except ContentParseError:
        raise
    except Exception as exc:
        raise ContentParseError(f"XLSX parsing failed: {exc}") from exc


def _parse_pptx(data: bytes) -> ParsedContent:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ContentParseError("PPTX parsing requires python-pptx") from exc
    try:
        presentation = Presentation(io.BytesIO(data))
        slides = []
        for number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            if parts:
                slides.append(f"[SLIDE {number}]\n" + "\n".join(parts))
        if not slides:
            raise ContentParseError("PPTX contains no extractable text")
        return ParsedContent(text="\n\n".join(slides), content_type="presentation", metadata={"format": "pptx", "slides": len(presentation.slides), "text_slides": len(slides)})
    except ContentParseError:
        raise
    except Exception as exc:
        raise ContentParseError(f"PPTX parsing failed: {exc}") from exc


def _parse_html(data: bytes) -> ParsedContent:
    from bs4 import BeautifulSoup
    html = _decode_text(data)
    soup = BeautifulSoup(html, "lxml")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    for tag in soup(["script", "style", "noscript", "nav", "footer", "aside", "form", "iframe"]):
        tag.decompose()
    text = soup.get_text(" ", strip=True)
    links = [a.get("href") for a in soup.find_all("a", href=True)]
    links = [link for link in links if link]
    if not text:
        raise ContentParseError("HTML contains no extractable text")
    return ParsedContent(text=text, title=title, content_type="html", links=links, metadata={"format": "html"})


def parse_content(url: str, content_type: str, data: bytes) -> ParsedContent:
    """Parse bytes according to MIME type/extension; never silently discard data."""
    kind = _kind(url, content_type)
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "csv":
        return _parse_csv(data)
    if kind == "tsv":
        return _parse_csv(data, "\t")
    if kind == "json":
        return _parse_json(data, False)
    if kind == "jsonl":
        return _parse_json(data, True)
    if kind == "xml":
        return _parse_xml(data)
    if kind == "docx":
        return _parse_docx(data)
    if kind == "xlsx":
        return _parse_xlsx(data)
    if kind == "pptx":
        return _parse_pptx(data)
    if kind == "html":
        return _parse_html(data)
    if kind == "text":
        text = _decode_text(data).strip()
        if not text:
            raise ContentParseError("text document is empty")
        return ParsedContent(text=text, content_type="text", metadata={"format": "text"})
    raise ContentParseError(f"unsupported content type: {content_type or 'unknown'} for {url}")
